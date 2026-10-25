#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""简历解析：文件 → 纯文本 + 质量判定。

输出结构固定为：
    {"raw_text": str, "parse_quality": "ok" | "partial" | "failed", "warnings": [str]}

设计要点
--------
1. 主路径用 restruct-cv 的 ingestion 层（native / docx / ocr），它是纯文本抽取，
   不触发 extract_resume 那两个 ONNX 模型，因此离线可用。
2. 任何一步失败都降级到 PyMuPDF / python-docx，绝不让单个库的问题变成整条链路失败。
3. 质量判定全部由代码规则给出，可复现、可解释；触发条件写进 warnings。
4. 少于 80 字符一律 failed。

命令行
------
    python v2/scripts/parse_resume.py 文件 [文件...] [--out 输出目录] [--pretty]
"""

from __future__ import annotations

import json
import re
import sys
import tempfile
from contextlib import contextmanager
from pathlib import Path

MIN_CHARS = 80
SPACE_DENSITY_LIMIT = 0.25   # 空格数 / 非空白字符数；中文 OCR 错乱时急剧升高
WEIRD_RATIO_LIMIT = 0.10     # 非预期字符占比
SECTION_KEYWORDS = ("教育", "学历", "毕业", "本科", "大专", "硕士", "博士",
                    "工作", "经验", "任职", "公司", "项目", "技能", "求职")

TMP_DIR = Path(__file__).resolve().parent.parent / ".tmp"

# 允许出现在文本里的字符：CJK、ASCII 字母数字、常见中英文标点、空白
_ALLOWED_PUNCT = set("，。、；：！？（）《》〈〉“”‘’—…·「」【】"
                     "_,.;:!?()[]%+/@#&*'\"~-—–$￥℃")
_CJK_RANGES = ((0x4E00, 0x9FFF), (0x3400, 0x4DBF), (0xF900, 0xFAFF))


# --------------------------------------------------------------------------- #
# 基础设施
# --------------------------------------------------------------------------- #
@contextmanager
def _noop_tempdir(tmp_path: Path):
    """restruct 的 native 层无条件创建临时目录做 OCR 兜底。

    本环境（以及很多受限容器）不允许 chmod/删除临时目录，TemporaryDirectory
    的清理会抛 PermissionError。这里把它换成"只给路径、不清理"的空壳，
    用完即恢复，避免污染全局状态。
    """
    tmp_path.mkdir(parents=True, exist_ok=True)
    original = tempfile.TemporaryDirectory

    class _Fixed:
        def __init__(self, *args, **kwargs):
            self.name = str(tmp_path)

        def __enter__(self):
            return self.name

        def __exit__(self, *exc):
            return False

    tempfile.TemporaryDirectory = _Fixed
    try:
        yield
    finally:
        tempfile.TemporaryDirectory = original


def _clean(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = "\n".join(line.rstrip() for line in text.split("\n"))
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


# --------------------------------------------------------------------------- #
# 各格式抽取
# --------------------------------------------------------------------------- #
def _extract_txt(path: Path, warnings: list[str]) -> str:
    for enc in ("utf-8-sig", "utf-8", "gb18030", "big5"):
        try:
            text = path.read_text(encoding=enc)
            warnings.append(f"文本文件按 {enc} 解码")
            return text
        except UnicodeDecodeError:
            continue
    raise ValueError("无法用 utf-8 / gb18030 / big5 解码该文本文件")


def _extract_pdf(path: Path, warnings: list[str]) -> str:
    try:
        import pymupdf
        from restruct.ingestion import native

        with _noop_tempdir(TMP_DIR):
            with pymupdf.open(path) as doc:
                pages = doc.page_count
                lines = native.extracted_lines(native.read_document(doc))
        used_ocr = any(getattr(l, "used_ocr", False) for l in lines)
        text = "\n".join(l.text for l in lines if l.text.strip())
        if used_ocr:
            warnings.append("部分页面使用了 OCR 识别，可能存在字符错误")
        if len(text.strip()) >= MIN_CHARS:
            warnings.append(f"PDF 由 restruct 抽取：{pages} 页 / {len(lines)} 行")
            return text
        warnings.append(f"restruct 仅抽出 {len(text.strip())} 字符（{pages} 页），改用 PyMuPDF 兜底")
    except Exception as e:  # noqa: BLE001 - 任何库内异常都不该中断链路
        warnings.append(f"restruct 抽取失败：{type(e).__name__}: {e}")

    import pymupdf

    with pymupdf.open(path) as doc:
        pages = doc.page_count
        text = "\n".join(page.get_text("text") for page in doc)
    warnings.append(f"PDF 由 PyMuPDF 抽取：{pages} 页")
    return text


def _extract_docx(path: Path, warnings: list[str]) -> str:
    try:
        from restruct.ingestion import docx as rdocx
        from restruct.ingestion import native

        document = rdocx.read_docx(path)
        lines = native.extracted_lines(document)
        text = "\n".join(l.text for l in lines if l.text.strip())
        if len(text.strip()) >= MIN_CHARS:
            warnings.append(f"DOCX 由 restruct 抽取：{len(lines)} 行")
            return text
        warnings.append(f"restruct 仅抽出 {len(text.strip())} 字符，改用 python-docx 兜底")
    except Exception as e:  # noqa: BLE001
        warnings.append(f"restruct 抽取失败：{type(e).__name__}: {e}")

    from docx import Document

    document = Document(str(path))
    parts = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))
    warnings.append("DOCX 由 python-docx 抽取（含表格）")
    return "\n".join(parts)


def _extract_image(path: Path, warnings: list[str]) -> str:
    from restruct.ingestion import ocr

    if not ocr.find_tesseract():
        raise RuntimeError(
            "图片/扫描件需要 OCR，但未在 PATH 中找到 Tesseract；"
            "请安装 tesseract 或改用可复制文本的 PDF/DOCX"
        )
    import pymupdf
    from restruct.ingestion import native

    with _noop_tempdir(TMP_DIR):
        with pymupdf.open(path) as doc:
            lines = native.extracted_lines(native.read_document(doc))
    warnings.append("图片由 Tesseract OCR 识别")
    return "\n".join(l.text for l in lines if l.text.strip())


# --------------------------------------------------------------------------- #
# 质量判定
# --------------------------------------------------------------------------- #
def _ok_char(ch: str) -> bool:
    if ch.isspace() or ch in _ALLOWED_PUNCT:
        return True
    if ch.isascii() and (ch.isalnum() or ch in "-+_.@/'\"()[]{}#&*:;,?!%$"):
        return True
    code = ord(ch)
    for lo, hi in _CJK_RANGES:
        if lo <= code <= hi:
            return True
    # 全角标点/字母数字
    if 0xFF00 <= code <= 0xFFEF:
        return ch in "！＂＃＄％＆＇（）＊＋，－．／：；＜＝＞？＠［＼］＾＿｀｛｜｝～" or ch.isalnum()
    return False


def classify(text: str, warnings: list[str]) -> str:
    """按确定性规则给出 ok / partial / failed，并把判定依据写进 warnings。"""
    stripped = text.strip()
    n = len(stripped)
    if n == 0:
        warnings.append("解析结果为 0 字符：文件可能为空、加密或为纯图片")
        return "failed"
    if n < MIN_CHARS:
        warnings.append(f"文本仅 {n} 字符，低于 {MIN_CHARS} 字符阈值，判定 failed")
        return "failed"

    body = stripped.replace(" ", "").replace("\u3000", "")
    spaces = stripped.count(" ") + stripped.count("\u3000")
    density = spaces / max(1, len(body))
    weird = sum(1 for ch in stripped if not _ok_char(ch))
    weird_ratio = weird / max(1, n)
    sections = [k for k in SECTION_KEYWORDS if k in stripped]

    if density > SPACE_DENSITY_LIMIT:
        warnings.append(f"空格密度 {density:.2f} > {SPACE_DENSITY_LIMIT}，疑似 OCR 断字或排版错乱")
        return "partial"
    if weird_ratio > WEIRD_RATIO_LIMIT:
        warnings.append(f"异常字符占比 {weird_ratio:.1%} > {WEIRD_RATIO_LIMIT:.0%}")
        return "partial"
    if not sections:
        warnings.append("未发现常见简历小节关键词（教育/工作/项目/技能等），内容可能不完整")
        return "partial"

    warnings.append(f"文本 {n} 字符，命中小节关键词 {len(sections)} 个，格式正常")
    return "ok"


# --------------------------------------------------------------------------- #
# 对外接口
# --------------------------------------------------------------------------- #
def parse_resume(path: str | Path) -> dict:
    """解析单个文件，返回 {raw_text, parse_quality, warnings}。"""
    p = Path(path)
    warnings: list[str] = []
    text = ""

    if not p.exists():
        return {"raw_text": "", "parse_quality": "failed",
                "warnings": [f"文件不存在：{p}"]}
    if not p.is_file():
        return {"raw_text": "", "parse_quality": "failed",
                "warnings": [f"不是文件：{p}"]}

    suffix = p.suffix.lower()
    try:
        if suffix in {".txt", ".md", ".text"}:
            text = _extract_txt(p, warnings)
        elif suffix == ".pdf":
            text = _extract_pdf(p, warnings)
        elif suffix in {".docx", ".doc"}:
            if suffix == ".doc":
                raise RuntimeError("旧版 .doc 不受支持，请另存为 .docx 或 PDF")
            text = _extract_docx(p, warnings)
        elif suffix in {".png", ".jpg", ".jpeg", ".bmp", ".tif", ".tiff", ".webp"}:
            text = _extract_image(p, warnings)
        else:
            raise RuntimeError(f"不支持的文件类型：{suffix or '（无扩展名）'}")
    except Exception as e:  # noqa: BLE001
        warnings.append(f"抽取失败：{type(e).__name__}: {e}")
        return {"raw_text": "", "parse_quality": "failed", "warnings": warnings}

    text = _clean(text)
    quality = classify(text, warnings)
    return {"raw_text": text, "parse_quality": quality, "warnings": warnings}


def main(argv: list[str]) -> int:
    import argparse

    ap = argparse.ArgumentParser(description="解析简历文件，输出 raw_text 与质量判定")
    ap.add_argument("paths", nargs="+")
    ap.add_argument("--out", default=None, help="把每个结果写成 <stem>.json 的目录")
    ap.add_argument("--pretty", action="store_true")
    args = ap.parse_args(argv)

    out_dir = Path(args.out) if args.out else None
    if out_dir:
        out_dir.mkdir(parents=True, exist_ok=True)

    for raw_path in args.paths:
        result = parse_resume(raw_path)
        if args.pretty:
            print(f"\n=== {raw_path} ===")
            print(f"quality: {result['parse_quality']}  chars: {len(result['raw_text'])}")
            for w in result["warnings"]:
                print(f"  - {w}")
        else:
            print(json.dumps(result, ensure_ascii=False))
        if out_dir:
            p = Path(raw_path)
            (out_dir / f"{p.stem}.json").write_text(
                json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    return 0


if __name__ == "__main__":
    raise SystemExit(main(sys.argv[1:]))
